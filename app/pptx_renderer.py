from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from app.schemas import PresentationPlan, SlidePlan


def _add_notes(slide, notes: str) -> None:
    slide.notes_slide.notes_text_frame.text = notes


def _layout(prs: Presentation, idx: int = 1):
    safe_idx = idx if len(prs.slide_layouts) > idx else 0
    return prs.slide_layouts[safe_idx]


def _render_slide(prs: Presentation, slide_plan: SlidePlan) -> None:
    slide = prs.slides.add_slide(_layout(prs, 1))
    slide.shapes.title.text = slide_plan.title

    body = None
    if len(slide.placeholders) > 1:
        body = slide.placeholders[1].text_frame
        body.clear()

    if slide_plan.type in {"title", "section"}:
        if body and slide_plan.subtitle:
            body.text = slide_plan.subtitle
    elif slide_plan.type in {"agenda", "bullets", "conclusion"}:
        if body:
            for i, bullet in enumerate(slide_plan.bullets):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = bullet
    elif slide_plan.type == "comparison":
        if body:
            left = ", ".join(slide_plan.left_bullets)
            right = ", ".join(slide_plan.right_bullets)
            body.text = f"{slide_plan.left_title}: {left}\n{slide_plan.right_title}: {right}"
    elif slide_plan.type == "quote":
        if body:
            body.text = slide_plan.subtitle or (slide_plan.bullets[0] if slide_plan.bullets else "")

    _add_notes(slide, slide_plan.speaker_notes)



def render_pptx(
    plan: PresentationPlan,
    output_path: str | Path,
    template_path: str | Path | None = None,
) -> None:
    prs = Presentation(str(template_path)) if template_path else Presentation()
    for slide_plan in plan.slides:
        _render_slide(prs, slide_plan)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
