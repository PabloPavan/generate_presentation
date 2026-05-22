from pathlib import Path

from pptx import Presentation

from app.pptx_renderer import render_pptx
from app.schemas import PresentationPlan


def _sample_plan() -> PresentationPlan:
    return PresentationPlan.model_validate(
        {
            "title": "Deck",
            "audience": "Alunos",
            "objective": "Ensinar",
            "slides": [
                {
                    "type": "title",
                    "title": "Intro",
                    "subtitle": "Sub",
                    "speaker_notes": "Notas intro",
                },
                {
                    "type": "conclusion",
                    "title": "Fim",
                    "bullets": ["Resumo"],
                    "speaker_notes": "Notas fim",
                },
            ],
        }
    )


def test_render_without_template(tmp_path: Path) -> None:
    out = tmp_path / "out.pptx"
    render_pptx(_sample_plan(), out)
    prs = Presentation(out)
    assert len(prs.slides) == 2


def test_render_with_template(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    Presentation().save(template)
    out = tmp_path / "with_template.pptx"
    render_pptx(_sample_plan(), out, template_path=template)
    prs = Presentation(out)
    assert len(prs.slides) == 2
